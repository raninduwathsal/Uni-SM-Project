import os
import re
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

def clean_markdown_text(content):
    # Encoding & typographical fixes
    content = content.replace('Welchs', "Welch's")
    content = content.replace('Welchs', "Welch's")
    content = content.replace('Cramr', 'Cramér')
    content = content.replace('Cramr', 'Cramér')
    content = content.replace('Nave', 'Naïve')
    content = content.replace('Nave', 'Naïve')

    # Specific math replacements
    replacements = [
        # Item 18
        (r'$t = 3.5810, p = 0.00034$', 't = 3.5810, p = 0.00034'),
        # Item 20
        (r'Seasonal $\text{ARIMA}(0,0,0)(0,1,1)[7]$ model', 'Seasonal ARIMA(0,0,0)(0,1,1)[7] model'),
        (r'($p = 0.398$)', '(p = 0.398)'),
        # Item 110 & 111
        (r'$\text{Revenue} = (\text{AdjustedRate} \times \text{Qty1}) - \text{Discount}$', '`Revenue = (AdjustedRate × Qty1) - Discount`'),
        (r'$\text{TotalBill} = \text{Revenue} + \text{ServiceCharge}$', '`TotalBill = Revenue + ServiceCharge`'),
        # Item 114 & 125
        (r'($1.5 \times \text{IQR}$)', '(1.5 × IQR)'),
        (r'($|z| > 3$)', '(|z| > 3)'),
        (r'($0.25\%$)', '(0.25%)'),
        (r'($>2.0$)', '(> 2.0)'),
        # Item 136
        (r'$\alpha = 0.05$', 'α = 0.05'),
        # Test 1
        (r'$H_0: \mu_{\text{Foreign}} - \mu_{\text{Local}} = 0 \quad \text{vs.} \quad H_1: \mu_{\text{Foreign}} - \mu_{\text{Local}} \neq 0$', 'H₀: μ_Foreign - μ_Local = 0   vs.   H₁: μ_Foreign - μ_Local ≠ 0'),
        (r'$\bar{x}_{\text{Foreign}} = \text{LKR } 1,039.90, \; \bar{x}_{\text{Local}} = \text{LKR } 1,014.57$', 'x̄_Foreign = LKR 1,039.90, x̄_Local = LKR 1,014.57'),
        (r'$t = 3.5810, \; df = 6,140.4, \; p = 0.00034$', 't = 3.5810, df = 6,140.4, p = 0.00034'),
        (r"Cohen's $d = 0.0594$", "Cohen's d = 0.0594"),
        (r'**Reject $H_0$.**', '**Reject H₀.**'),
        # Test 2
        (r'$H_0: \sigma^2_{\text{Dine-In}} / \sigma^2_{\text{Takeaway}} = 1 \quad \text{vs.} \quad H_1: \sigma^2_{\text{Dine-In}} / \sigma^2_{\text{Takeaway}} \neq 1$', 'H₀: σ²_Dine-In / σ²_Takeaway = 1   vs.   H₁: σ²_Dine-In / σ²_Takeaway ≠ 1'),
        (r'$F = 0.9687, \; df_1 = 17,364, \; df_2 = 12,099, \; p = 0.05723$', 'F = 0.9687, df₁ = 17,364, df₂ = 12,099, p = 0.05723'),
        (r'**Fail to Reject $H_0$.**', '**Fail to Reject H₀.**'),
        # Test 3
        (r'$H_0: \mu_1 = \mu_2 = \dots = \mu_8 \quad \text{vs.} \quad H_1: \text{At least one outlet mean differs}$', 'H₀: μ₁ = μ₂ = ... = μ₈   vs.   H₁: At least one outlet mean differs'),
        (r'$F = 1.4044, \; df = (7, 5745), \; p = 0.1986, \; \eta^2 = 0.0017$', 'F = 1.4044, df = (7, 5745), p = 0.1986, η² = 0.0017'),
        (r'With $p > 0.05$,', 'With p > 0.05,'),
        # Test 4
        (r'$H_0: \text{Payment Method} \perp \text{Customer Type} \quad \text{vs.} \quad H_1: \text{Statistically Dependent}$', 'H₀: Payment Method ⊥ Customer Type   vs.   H₁: Statistically Dependent'),
        (r'$\chi^2 = 4,270.2, \; df = 6, \; p < 10^{-4}$', 'χ² = 4,270.2, df = 6, p < 0.0001'),
        (r"Cramér's $V = 0.3807$", "Cramér's V = 0.3807"),
        (r"Cramr's $V = 0.3807$", "Cramér's V = 0.3807"),
        # Test 5
        (r'$H_0: \text{Medians across Morning, Afternoon, Evening, Night identical}$', 'H₀: Medians across Morning, Afternoon, Evening, Night identical'),
        (r'Kruskal-Wallis $\chi^2 = 0.7777, \; df = 3, \; p = 0.8548$', 'Kruskal-Wallis χ² = 0.7777, df = 3, p = 0.8548'),
        # Test 6
        (r'$H_0: \rho = 0 \quad \text{vs.} \quad H_1: \rho \neq 0$', 'H₀: ρ = 0   vs.   H₁: ρ ≠ 0'),
        (r'Pearson $r = +0.0336, \; t = 5.7624, \; df = 29,463, \; p < 10^{-4}$', 'Pearson r = +0.0336, t = 5.7624, df = 29,463, p < 0.0001'),
        (r'Spearman $\rho = +0.0335, \; p < 10^{-4}$', 'Spearman ρ = +0.0335, p < 0.0001'),
        # Task 6 Model Block
        (r'$$\text{Model: } Y_{ij} = \mu + \tau_i (\text{Discount}) + \beta_j (\text{Store Block}) + \varepsilon_{ij}$$', '> **RCBD Model Specification:**  \n> `Y_ij = μ + τ_i(Discount) + β_j(Store Block) + ε_ij`'),
        # Task 7
        (r'($p = 5$)', '(p = 5)'),
        (r'($n = 29,465$)', '(n = 29,465)'),
        # Task 8
        (r'$P(\text{Beverage}) = 53.91\%, \; P(\text{Food}) = 37.26\%, \; P(\text{Deals}) = 4.60\%, \; P(\text{Merchandizing}) = 4.23\%$', 'P(Beverage) = 53.91%, P(Food) = 37.26%, P(Deals) = 4.60%, P(Merchandizing) = 4.23%'),
        (r'($\alpha = 1$)', '(α = 1)'),
        (r'$L(y, \hat{y})$', 'L(y, ŷ)'),
        # Task 9
        (r'($p < 0.001$)', '(p < 0.001)'),
        (r'$$\text{Model: } \text{ARIMA}(0,0,0)(0,1,1)_7 \text{ with drift}$$', '> **Selected Model:**  \n> `Seasonal ARIMA(0,0,0)(0,1,1)[7] with drift`'),
        (r'$\text{sma1} = -0.9201 \; (se = 0.2101)$', 'sma1 = -0.9201 (se = 0.2101)'),
        (r'$\text{drift} = -0.3441$', 'drift = -0.3441'),
        (r'$\chi^2 = 14.7103, \; df = 14, \; p = 0.39824$', 'χ² = 14.7103, df = 14, p = 0.39824'),
        (r'Because $p > 0.05$,', 'Because p > 0.05,'),
        # Task 12
        (r'$1.96 \times \text{RMSE}$', '1.96 × RMSE'),
    ]

    for old, new in replacements:
        content = content.replace(old, new)
        
    # Catch any remaining single dollar math artifacts if any
    content = re.sub(r'\$(.*?)\$', r'\1', content)
    return content

def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def add_inline_runs(paragraph, text, default_color=None):
    # Regex splits on bold (**...**), italic (*...*), and inline code (`...`)
    parts = re.split(r'(\*\*.*?\*\*|\*.*?\*|`.*?`)', text)
    for part in parts:
        if not part:
            continue
        if part.startswith('**') and part.endswith('**'):
            r = paragraph.add_run(part[2:-2])
            r.font.bold = True
            r.font.name = 'Calibri'
            if default_color:
                r.font.color.rgb = default_color
        elif part.startswith('*') and part.endswith('*'):
            r = paragraph.add_run(part[1:-1])
            r.font.italic = True
            r.font.name = 'Calibri'
            if default_color:
                r.font.color.rgb = default_color
        elif part.startswith('`') and part.endswith('`'):
            r = paragraph.add_run(part[1:-1])
            r.font.name = 'Consolas'
            r.font.size = Pt(10)
            r.font.color.rgb = RGBColor(160, 40, 60) # Deep crimson for code
        else:
            r = paragraph.add_run(part)
            r.font.name = 'Calibri'
            if default_color:
                r.font.color.rgb = default_color

def create_styled_docx(md_path, docx_path):
    with open(md_path, 'r', encoding='utf-8') as f:
        raw_text = f.read()
    
    clean_text = clean_markdown_text(raw_text)
    
    # Also save clean markdown back
    with open(md_path, 'w', encoding='utf-8') as f:
        f.write(clean_text)
        
    doc = docx.Document()
    
    # Standard 1-inch margins
    for section in doc.sections:
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)
        
    # Set default style
    style = doc.styles['Normal']
    font = style.font
    font.name = 'Calibri'
    font.size = Pt(11)
    font.color.rgb = RGBColor(40, 40, 40)
    
    lines = clean_text.splitlines()
    i = 0
    in_code_block = False
    code_block_lines = []
    table_rows = []
    
    def process_table_block(rows):
        if not rows:
            return
        parsed = []
        for r in rows:
            cells = [c.strip() for c in r.strip().strip('|').split('|')]
            if all(re.match(r'^:?-+:?$', c) for c in cells):
                continue
            parsed.append(cells)
        if not parsed:
            return
        
        col_count = max(len(r) for r in parsed)
        tbl = doc.add_table(rows=len(parsed), cols=col_count)
        tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
        
        for r_idx, row_data in enumerate(parsed):
            row = tbl.rows[r_idx]
            is_header = (r_idx == 0)
            for c_idx in range(col_count):
                cell = row.cells[c_idx]
                val = row_data[c_idx] if c_idx < len(row_data) else ""
                val_clean = re.sub(r'\*\*(.*?)\*\*', r'\1', val)
                val_clean = re.sub(r'\*(.*?)\*', r'\1', val_clean)
                val_clean = re.sub(r'`(.*?)`', r'\1', val_clean)
                
                p = cell.paragraphs[0]
                p.text = val_clean
                p.paragraph_format.space_before = Pt(3)
                p.paragraph_format.space_after = Pt(3)
                set_cell_margins(cell, top=120, bottom=120, left=140, right=140)
                
                if is_header:
                    set_cell_background(cell, "1F4E78") # Navy
                    if p.runs:
                        p.runs[0].font.bold = True
                        p.runs[0].font.color.rgb = RGBColor(255, 255, 255)
                        p.runs[0].font.size = Pt(10)
                else:
                    if p.runs:
                        p.runs[0].font.size = Pt(9.5)
                    if r_idx % 2 == 1:
                        set_cell_background(cell, "F2F5F8") # Subtle blue-grey alternate
                    else:
                        set_cell_background(cell, "FFFFFF")
        doc.add_paragraph() # Spacer
    
    while i < len(lines):
        line = lines[i]
        stripped = line.strip()
        
        # Check code block fences
        if stripped.startswith('```'):
            if in_code_block:
                # End of code block
                tbl = doc.add_table(rows=1, cols=1)
                tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
                cell = tbl.rows[0].cells[0]
                set_cell_background(cell, "F4F6F8")
                set_cell_margins(cell, top=140, bottom=140, left=180, right=180)
                cp = cell.paragraphs[0]
                cp.text = "\n".join(code_block_lines)
                cp.runs[0].font.name = 'Consolas'
                cp.runs[0].font.size = Pt(9.5)
                cp.runs[0].font.color.rgb = RGBColor(50, 50, 50)
                doc.add_paragraph() # Spacer
                code_block_lines = []
                in_code_block = False
            else:
                in_code_block = True
            i += 1
            continue
            
        if in_code_block:
            code_block_lines.append(line)
            i += 1
            continue
            
        # Check table
        if stripped.startswith('|') and '|' in stripped[1:]:
            table_rows.append(stripped)
            i += 1
            continue
        else:
            if table_rows:
                process_table_block(table_rows)
                table_rows = []
                
        # Empty line
        if not stripped:
            i += 1
            continue
            
        # Horizontal rule
        if stripped in ['---', '***', '___']:
            i += 1
            continue
            
        # Headings
        if stripped.startswith('# '):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(20)
            p.paragraph_format.space_after = Pt(8)
            run = p.add_run(stripped[2:].strip())
            run.font.name = 'Calibri'
            run.font.size = Pt(22)
            run.font.bold = True
            run.font.color.rgb = RGBColor(31, 78, 120)
        elif stripped.startswith('## '):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(16)
            p.paragraph_format.space_after = Pt(6)
            run = p.add_run(stripped[3:].strip())
            run.font.name = 'Calibri'
            run.font.size = Pt(16)
            run.font.bold = True
            run.font.color.rgb = RGBColor(31, 78, 120)
        elif stripped.startswith('### '):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(12)
            p.paragraph_format.space_after = Pt(4)
            run = p.add_run(stripped[4:].strip())
            run.font.name = 'Calibri'
            run.font.size = Pt(13)
            run.font.bold = True
            run.font.color.rgb = RGBColor(50, 60, 80)
        elif stripped.startswith('#### '):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(9)
            p.paragraph_format.space_after = Pt(3)
            run = p.add_run(stripped[5:].strip())
            run.font.name = 'Calibri'
            run.font.size = Pt(11.5)
            run.font.bold = True
            run.font.color.rgb = RGBColor(70, 80, 95)
        # Bullet points
        elif stripped.startswith('- ') or stripped.startswith('* '):
            text = stripped[2:].strip()
            p = doc.add_paragraph(style='List Bullet')
            p.paragraph_format.space_before = Pt(1.5)
            p.paragraph_format.space_after = Pt(2.5)
            add_inline_runs(p, text)
        # Numbered list
        elif re.match(r'^\d+\.\s', stripped):
            text = re.sub(r'^\d+\.\s', '', stripped).strip()
            p = doc.add_paragraph(style='List Number')
            p.paragraph_format.space_before = Pt(1.5)
            p.paragraph_format.space_after = Pt(2.5)
            add_inline_runs(p, text)
        # Blockquote / Callout
        elif stripped.startswith('> '):
            text = stripped[2:].strip()
            tbl = doc.add_table(rows=1, cols=1)
            tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
            cell = tbl.rows[0].cells[0]
            set_cell_background(cell, "F0F4F8") # Elegant light tint
            set_cell_margins(cell, top=100, bottom=100, left=160, right=160)
            p = cell.paragraphs[0]
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(2)
            add_inline_runs(p, text, default_color=RGBColor(45, 60, 85))
            doc.add_paragraph() # Spacer
        # Standard paragraph
        else:
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(4)
            p.paragraph_format.line_spacing = 1.15
            add_inline_runs(p, stripped)
        i += 1
        
    if table_rows:
        process_table_block(table_rows)
        
    # Save target
    os.makedirs(os.path.dirname(docx_path), exist_ok=True)
    doc.save(docx_path)
    print(f"[SUCCESS] Saved Word document: {docx_path}")

if __name__ == "__main__":
    files = [
        ("FINAL/01_Reports_and_Documentation/IT3081_Final_Consultancy_Report.md",
         "FINAL/01_Reports_and_Documentation/IT3081_Final_Consultancy_Report.docx"),
        ("Courseweb Documents/IT3081_Final_Consultancy_Report.md",
         "Courseweb Documents/IT3081_Final_Consultancy_Report.docx")
    ]
    for md, docx_f in files:
        create_styled_docx(md, docx_f)
