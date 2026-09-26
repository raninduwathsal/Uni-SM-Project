import os
import re
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

md_path = "Courseweb Documents/IT3081_Final_Consultancy_Report.md"
docx_path = "Courseweb Documents/IT3081_Final_Consultancy_Report.docx"

doc = docx.Document()

# Set standard 1-inch margins
for section in doc.sections:
    section.top_margin = Inches(1.0)
    section.bottom_margin = Inches(1.0)
    section.left_margin = Inches(1.0)
    section.right_margin = Inches(1.0)

# Set Normal Style font
style = doc.styles['Normal']
font = style.font
font.name = 'Calibri'
font.size = Pt(11)
font.color.rgb = RGBColor(40, 40, 40)

def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

with open(md_path, 'r', encoding='utf-8') as f:
    lines = f.readlines()

in_table = False
table_rows = []

def process_table(rows):
    if not rows:
        return
    # Parse rows
    parsed = []
    for r in rows:
        cells = [c.strip() for c in r.strip().strip('|').split('|')]
        # Skip separator row like |---|---|
        if all(re.match(r'^:?-+:?$', c) for c in cells):
            continue
        parsed.append(cells)
    
    if not parsed:
        return
    
    col_count = max(len(r) for r in parsed)
    tbl = doc.add_table(rows=len(parsed), cols=col_count)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    
    for r_idx, row_data in enumerate(parsed):
        row = tbl.rows[r_idx]
        is_header = (r_idx == 0)
        for c_idx in range(col_count):
            cell = row.cells[c_idx]
            val = row_data[c_idx] if c_idx < len(row_data) else ""
            # Clean bold markers
            val_clean = re.sub(r'\*\*(.*?)\*\*', r'\1', val)
            val_clean = re.sub(r'\*(.*?)\*', r'\1', val_clean)
            val_clean = re.sub(r'`(.*?)`', r'\1', val_clean)
            val_clean = re.sub(r'\$(.*?)\$', r'\1', val_clean)
            
            p = cell.paragraphs[0]
            p.text = val_clean
            p.paragraph_format.space_before = Pt(3)
            p.paragraph_format.space_after = Pt(3)
            set_cell_margins(cell, top=120, bottom=120, left=140, right=140)
            
            if is_header:
                set_cell_background(cell, "1F4E78")
                p.runs[0].font.bold = True
                p.runs[0].font.color.rgb = RGBColor(255, 255, 255)
                p.runs[0].font.size = Pt(10)
            else:
                p.runs[0].font.size = Pt(9.5)
                if r_idx % 2 == 1:
                    set_cell_background(cell, "F2F5F8")
                else:
                    set_cell_background(cell, "FFFFFF")
    
    doc.add_paragraph() # Spacer

i = 0
while i < len(lines):
    line = lines[i].rstrip('\n')
    stripped = line.strip()
    
    # Check table
    if stripped.startswith('|') and '|' in stripped[1:]:
        table_rows.append(stripped)
        i += 1
        continue
    else:
        if table_rows:
            process_table(table_rows)
            table_rows = []
    
    # Empty line
    if not stripped:
        i += 1
        continue
    
    # Horizontal rule
    if stripped in ['---', '***', '___']:
        i += 1
        continue
    
    # Headings
    if stripped.startswith('# '):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(16)
        p.paragraph_format.space_after = Pt(6)
        run = p.add_run(stripped[2:].strip())
        run.font.name = 'Calibri'
        run.font.size = Pt(22)
        run.font.bold = True
        run.font.color.rgb = RGBColor(31, 78, 120)
    elif stripped.startswith('## '):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(14)
        p.paragraph_format.space_after = Pt(5)
        run = p.add_run(stripped[3:].strip())
        run.font.name = 'Calibri'
        run.font.size = Pt(16)
        run.font.bold = True
        run.font.color.rgb = RGBColor(31, 78, 120)
    elif stripped.startswith('### '):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(10)
        p.paragraph_format.space_after = Pt(4)
        run = p.add_run(stripped[4:].strip())
        run.font.name = 'Calibri'
        run.font.size = Pt(13)
        run.font.bold = True
        run.font.color.rgb = RGBColor(50, 60, 80)
    elif stripped.startswith('#### '):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(8)
        p.paragraph_format.space_after = Pt(2)
        run = p.add_run(stripped[5:].strip())
        run.font.name = 'Calibri'
        run.font.size = Pt(11.5)
        run.font.bold = True
        run.font.color.rgb = RGBColor(70, 80, 95)
    # Bullet points
    elif stripped.startswith('- ') or stripped.startswith('* '):
        text = stripped[2:].strip()
        p = doc.add_paragraph(style='List Bullet')
        p.paragraph_format.space_before = Pt(1)
        p.paragraph_format.space_after = Pt(2)
        # Parse inline formatting
        parts = re.split(r'(\*\*.*?\*\*|\*.*?\*|`.*?`)', text)
        for part in parts:
            if not part: continue
            if part.startswith('**') and part.endswith('**'):
                r = p.add_run(part[2:-2])
                r.font.bold = True
            elif part.startswith('*') and part.endswith('*'):
                r = p.add_run(part[1:-1])
                r.font.italic = True
            elif part.startswith('`') and part.endswith('`'):
                r = p.add_run(part[1:-1])
                r.font.name = 'Consolas'
                r.font.size = Pt(10)
            else:
                p.add_run(part)
    # Numbered list
    elif re.match(r'^\d+\.\s', stripped):
        text = re.sub(r'^\d+\.\s', '', stripped).strip()
        p = doc.add_paragraph(style='List Number')
        p.paragraph_format.space_before = Pt(1)
        p.paragraph_format.space_after = Pt(2)
        parts = re.split(r'(\*\*.*?\*\*|\*.*?\*|`.*?`)', text)
        for part in parts:
            if not part: continue
            if part.startswith('**') and part.endswith('**'):
                r = p.add_run(part[2:-2])
                r.font.bold = True
            elif part.startswith('*') and part.endswith('*'):
                r = p.add_run(part[1:-1])
                r.font.italic = True
            elif part.startswith('`') and part.endswith('`'):
                r = p.add_run(part[1:-1])
                r.font.name = 'Consolas'
                r.font.size = Pt(10)
            else:
                p.add_run(part)
    # Blockquote
    elif stripped.startswith('> '):
        text = stripped[2:].strip()
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Inches(0.4)
        p.paragraph_format.space_before = Pt(4)
        p.paragraph_format.space_after = Pt(4)
        r = p.add_run(text)
        r.font.italic = True
        r.font.color.rgb = RGBColor(80, 90, 105)
    # Standard text
    else:
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(2)
        p.paragraph_format.space_after = Pt(4)
        parts = re.split(r'(\*\*.*?\*\*|\*.*?\*|`.*?`)', stripped)
        for part in parts:
            if not part: continue
            if part.startswith('**') and part.endswith('**'):
                r = p.add_run(part[2:-2])
                r.font.bold = True
            elif part.startswith('*') and part.endswith('*'):
                r = p.add_run(part[1:-1])
                r.font.italic = True
            elif part.startswith('`') and part.endswith('`'):
                r = p.add_run(part[1:-1])
                r.font.name = 'Consolas'
                r.font.size = Pt(10)
            else:
                p.add_run(part)
    i += 1

if table_rows:
    process_table(table_rows)

doc.save(docx_path)
print(f"[SUCCESS] Successfully created Word document: {docx_path}")
